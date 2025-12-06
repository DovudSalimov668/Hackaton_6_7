"""
Views для финансовых инструментов
"""
import json
from django.shortcuts import render
from django.views.generic import TemplateView
from . import services
from .models import PaymentSchedule
from django_filters.views import FilterView
import django_filters
from crispy_forms.helper import FormHelper
from crispy_forms.layout import Submit
from django import forms
from django.http import HttpResponse
import openpyxl
from tempfile import NamedTemporaryFile


# Example: Filter form for expense analysis
class ExpensesCategoryFilterForm(forms.Form):
    query = forms.CharField(label="Поиск по категории", required=False)
    min_amount = forms.DecimalField(label="Минимальная сумма", required=False, decimal_places=2, min_value=0)
    max_amount = forms.DecimalField(label="Максимальная сумма", required=False, decimal_places=2, min_value=0)
    start_date = forms.DateField(label="С даты", required=False, widget=forms.DateInput(attrs={'type':'date'}))
    end_date = forms.DateField(label="По дату", required=False, widget=forms.DateInput(attrs={'type':'date'}))

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.helper = FormHelper()
        self.helper.form_method = 'get'
        self.helper.form_class = 'row g-3 mb-4'
        self.helper.label_class = 'form-label col-form-label col-auto'
        self.helper.add_input(Submit('submit', 'Фильтровать', css_class='btn btn-primary'))


class ExpensesView(TemplateView):
    """Кейс 4: Анализ расходов"""
    template_name = 'expenses.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Prepare the filter form
        form = ExpensesCategoryFilterForm(self.request.GET or None)
        context['filter_form'] = form

        # Получаем данные анализа расходов
        analysis = services.analyze_expenses()

        context['total_expenses'] = analysis['total']
        context['categories'] = analysis['categories']

        # Данные для Pie Chart (в JSON для JavaScript)
        context['chart_data'] = json.dumps({
            'labels': analysis['labels'],
            'data': analysis['data'],
            'colors': analysis['colors'],
        }, ensure_ascii=False)

        return context


class PaymentsView(TemplateView):
    """Кейс 5: Напоминания о платежах"""
    template_name = 'payments.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Получаем ближайший платеж
        upcoming = services.get_upcoming_payment()
        context['upcoming'] = upcoming

        # Получаем просроченные платежи
        overdue = services.get_overdue_payments()
        context['overdue'] = overdue

        # Весь график платежей
        context['all_payments'] = PaymentSchedule.objects.all()

        return context


class CalculatorView(TemplateView):
    """Кейс 6: Кредитный калькулятор"""
    template_name = 'calculator.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Значения по умолчанию
        context['amount'] = 100000
        context['months'] = 12
        context['rate'] = 15.0
        context['one_time_commission'] = 0
        context['monthly_commission'] = 0

        return context

    def post(self, request, *args, **kwargs):
        context = self.get_context_data(**kwargs)

        # Получаем параметры из формы
        try:
            amount = float(request.POST.get('amount', 100000))
            months = int(request.POST.get('months', 12))
            rate = float(request.POST.get('rate', 15.0))
            one_time_commission = float(request.POST.get('one_time_commission', 0))
            monthly_commission = float(request.POST.get('monthly_commission', 0))

            # Валидация
            amount = max(100, min(500000, amount))
            months = max(1, min(120, months))
            rate = max(1, min(35, rate))
            one_time_commission = max(0, min(5, one_time_commission))
            monthly_commission = max(0, min(2, monthly_commission))

            context['amount'] = amount
            context['months'] = months
            context['rate'] = rate
            context['one_time_commission'] = one_time_commission
            context['monthly_commission'] = monthly_commission

            # Расчет обоих типов кредита с комиссиями
            comparison = services.calculate_loan_comparison(
                amount, months, rate, one_time_commission, monthly_commission
            )

            context['annuity_schedule'] = comparison['annuity']
            context['diff_schedule'] = comparison['differentiated']
            context['annuity_total'] = comparison['annuity_total']
            context['annuity_overpayment'] = comparison['annuity_overpayment']
            context['diff_total'] = comparison['diff_total']
            context['diff_overpayment'] = comparison['diff_overpayment']

            # Данные для графиков
            # 1. Stacked Bar Chart - платежи по месяцам для аннуитета
            annuity_labels = [f"Мес. {p['month']}" for p in comparison['annuity']]
            annuity_principal = [p['principal'] for p in comparison['annuity']]
            annuity_interest = [p['interest'] for p in comparison['annuity']]

            context['annuity_chart_data'] = json.dumps({
                'labels': annuity_labels,
                'principal': annuity_principal,
                'interest': annuity_interest,
            }, ensure_ascii=False)

            # 2. Stacked Bar Chart - платежи по месяцам для дифференцированного
            diff_labels = [f"Мес. {p['month']}" for p in comparison['differentiated']]
            diff_principal = [p['principal'] for p in comparison['differentiated']]
            diff_interest = [p['interest'] for p in comparison['differentiated']]

            context['diff_chart_data'] = json.dumps({
                'labels': diff_labels,
                'principal': diff_principal,
                'interest': diff_interest,
            }, ensure_ascii=False)

            # 3. Сравнительный Line Chart - общие платежи
            comparison_payments_annuity = [p['payment'] for p in comparison['annuity']]
            comparison_payments_diff = [p['payment'] for p in comparison['differentiated']]

            context['comparison_chart_data'] = json.dumps({
                'labels': annuity_labels,
                'annuity': comparison_payments_annuity,
                'differentiated': comparison_payments_diff,
            }, ensure_ascii=False)

            # 4. Pie Chart данные - тело долга vs переплата (для аннуитета)
            context['pie_chart_data'] = json.dumps({
                'labels': ['Тело кредита', 'Проценты', 'Комиссии'],
                'data': [
                    comparison['original_amount'],
                    comparison['annuity_interest'],
                    comparison['annuity_one_time_commission'] + comparison['annuity_monthly_commission']
                ],
                'colors': ['#00a651', '#ef4444', '#ffa500']
            }, ensure_ascii=False)

            # 5. Line Chart - остаток долга
            annuity_remaining = [p['remaining'] for p in comparison['annuity']]
            diff_remaining = [p['remaining'] for p in comparison['differentiated']]

            context['remaining_debt_chart_data'] = json.dumps({
                'labels': annuity_labels,
                'annuity': annuity_remaining,
                'differentiated': diff_remaining,
            }, ensure_ascii=False)

        except (ValueError, TypeError) as e:
            context['error'] = 'Ошибка в параметрах расчета'

        return render(request, self.template_name, context)



class CurrencyView(TemplateView):
    """Кейс 7: Обмен валют"""
    template_name = 'currency.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Значения по умолчанию
        context['amount'] = 100
        context['currency'] = 'USD'
        context['operation'] = 'buy'
        context['history_days'] = 30

        try:
            # Прогноз курса USD
            usd_forecast = services.get_currency_forecast('USD', days=7)
            context['usd_forecast'] = usd_forecast if usd_forecast else {}

            # Прогноз курса EUR
            eur_forecast = services.get_currency_forecast('EUR', days=7)
            context['eur_forecast'] = eur_forecast if eur_forecast else {}

            # История курсов для графика
            # Всегда возвращаем валидный JSON, даже если данных нет

            # USD
            try:
                usd_hist_data = services.get_currency_history('USD', days=30)
                if not usd_hist_data or not usd_hist_data.get('labels'):
                    usd_hist_data = {'labels': [], 'buy_data': [], 'sell_data': []}
            except Exception:
                usd_hist_data = {'labels': [], 'buy_data': [], 'sell_data': []}

            context['usd_history'] = json.dumps(usd_hist_data, ensure_ascii=False)

            # EUR
            try:
                eur_hist_data = services.get_currency_history('EUR', days=30)
                if not eur_hist_data or not eur_hist_data.get('labels'):
                    eur_hist_data = {'labels': [], 'buy_data': [], 'sell_data': []}
            except Exception:
                eur_hist_data = {'labels': [], 'buy_data': [], 'sell_data': []}

            context['eur_history'] = json.dumps(eur_hist_data, ensure_ascii=False)

        except Exception as e:
            print(f"Error in CurrencyView context: {e}")
            # Fallback values
            context['usd_forecast'] = {}
            context['eur_forecast'] = {}
            context['usd_history'] = '{}'
            context['eur_history'] = '{}'

        return context

    def post(self, request, *args, **kwargs):
        context = self.get_context_data(**kwargs)

        # Конвертация валюты
        try:
            amount = float(request.POST.get('amount', 100))
            currency = request.POST.get('currency', 'USD')
            operation = request.POST.get('operation', 'buy')
            history_days = int(request.POST.get('history_days', 30))

            context['amount'] = amount
            context['currency'] = currency
            context['operation'] = operation
            context['history_days'] = history_days

            # Выполняем конвертацию
            try:
                conversion = services.convert_currency(amount, currency, operation)
                context['conversion'] = conversion
            except Exception as e:
                context['error'] = f'Ошибка конвертации: {str(e)}'

        except (ValueError, TypeError) as e:
            context['error'] = 'Ошибка в параметрах конвертации'

        return render(request, self.template_name, context)


def export_expenses_excel(request):
    # Get expenses data (ideally reuse chart context)
    from . import services
    analysis = services.analyze_expenses()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Expenses Report'
    ws.append(["Категория", "Сумма", "%"])
    for cat in analysis['categories']:
        ws.append([cat['name'], cat['amount'], cat['percentage']])
    ws.append([])
    ws.append(["Итого", analysis['total'], 100])

    # Styling can be added here

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=expenses_report.xlsx'
    wb.save(response)
    return response


def export_calculator_excel(request):
    from . import services
    import openpyxl
    if request.method == "POST":
        # Default sample params, ideally parse user's form params or session
        amount = float(request.POST.get("amount", 100000))
        months = int(request.POST.get("months", 12))
        rate = float(request.POST.get("rate", 15))
        one_time_comm = float(request.POST.get("one_time_commission", 0))
        monthly_comm = float(request.POST.get("monthly_commission", 0))
        schedule_type = request.POST.get("schedule_type", "annuity")
        # Calculate corresponding schedule
        if schedule_type == "differentiated":
            schedule, _ = services.calculate_differentiated_loan(amount, months, rate, one_time_comm, monthly_comm)
        else:
            schedule, _ = services.calculate_annuity_loan(amount, months, rate, one_time_comm, monthly_comm)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Loan Schedule"
        ws.append(["Месяц", "Платеж", "Тело", "Проценты", "Комиссия", "Остаток"])
        for p in schedule:
            ws.append([
                p['month'], p['payment'], p['principal'], p['interest'], p.get('monthly_commission', 0), p.get('remaining', 0)
            ])
        from django.http import HttpResponse
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=loan_schedule.xlsx'
        wb.save(response)
        return response
    else:
        from django.http import HttpResponseBadRequest
        return HttpResponseBadRequest("Invalid method")


def export_payments_excel(request):
    from .models import PaymentSchedule
    import openpyxl
    if request.method == "POST":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Payments"
        ws.append(["Дата платежа", "Описание", "Запланировано", "Оплачено", "Статус"])
        for payment in PaymentSchedule.objects.order_by('due_date'):
            status = (
                "Оплачено" if payment.is_paid else ("Просрочка" if payment.is_overdue else "Не оплачено")
            )
            ws.append([
                str(payment.due_date), payment.description, float(payment.planned_amount), float(payment.paid_amount), status
            ])
        from django.http import HttpResponse
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=payments_schedule.xlsx'
        wb.save(response)
        return response
    else:
        from django.http import HttpResponseBadRequest
        return HttpResponseBadRequest("Invalid method")


class ExpensesFilter(django_filters.FilterSet):
    name = django_filters.CharFilter(field_name='category__name', lookup_expr='icontains', label='Категория')

    class Meta:
        model = None  # Non-model example for now
        fields = ['name']

# Optionally, inherit FilterView if you wish to add to class-based view (requires real queryset)
